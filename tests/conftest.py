import json
import re
import shutil
import subprocess
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest

HAS_FFMPEG = shutil.which("ffmpeg") is not None
needs_ffmpeg = pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg is not installed")

# How many contributions a part was asked for, in both languages of the prompt.
_COUNT = re.compile(r"(?:exactly (\d+) contributions|genau (\d+) Beiträge)")
_PARTS = re.compile(r"exactly (\d+) consecutive parts")
_TOPICS = re.compile(r"at most (\d+), spread over the whole text")


def _mp3(seconds=0.3) -> bytes:
    """A short real MP3, so ffmpeg can read and join it."""
    if not HAS_FFMPEG:
        return b"ID3fake-mp3"
    with tempfile.TemporaryDirectory() as tmp:
        target = Path(tmp) / "a.mp3"
        subprocess.run(
            ["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i",
             "sine=frequency=440:duration=%s" % seconds,
             "-ar", "24000", "-ac", "1", str(target)],
            check=True,
        )
        return target.read_bytes()


MP3 = _mp3()


def plan_answer(parts: int, files=None) -> str:
    """A plan with as many parts as asked for, each naming the files it draws on."""
    out = []
    for i in range(parts):
        part = {"title": "Part %d" % (i + 1), "gist": "What part %d covers." % (i + 1)}
        if files is not None:
            part["files"] = list(files[i]) if i < len(files) else []
        out.append(part)
    return json.dumps({
        "opening": "Lea stands in front of the notice board and wonders what to do.",
        "thread": "One rule, many everyday questions.",
        "parts": out,
    })


def topics_answer(most: int) -> str:
    return json.dumps({"topics": [{"title": "Topic %d" % (i + 1),
                                   "gist": "What topic %d states." % (i + 1)}
                                  for i in range(most)]})


def lines_answer(count: int, persons: int = 2) -> str:
    return json.dumps({
        "lines": [{"speaker": i % max(1, persons),
                   "long": i % 3 == 0,
                   "text": "Contribution number %d says something about the material." % (i + 1)}
                  for i in range(count)]
    })


class FakeServer:
    """Answers chat completions with a plan or a script, and speech with a short MP3."""

    def __init__(self):
        self.requests = []
        self.fail = False
        self.fail_status = 500
        self.refuse_json_mode = False
        self.speech_failures = 0
        self.speech_status = 500
        self.speech_delay = 0.0
        self.persons = 2
        # One entry per planning call: the files each part of that plan names.
        self.plan_files = []
        self.lock = threading.Lock()
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                length = int(self.headers.get("Content-Length", 0))
                body = json.loads(self.rfile.read(length) or b"{}")
                with outer.lock:
                    outer.requests.append({
                        "path": self.path,
                        "headers": {k.lower(): v for k, v in self.headers.items()},
                        "body": body,
                    })
                if outer.fail:
                    return self.refuse(outer.fail_status, "broken")
                if self.path.endswith("/audio/speech"):
                    return self.speech(body)
                if self.path.endswith("/chat/completions"):
                    return self.chat(body)
                self.refuse(404, "unknown path")

            def speech(self, body):
                with outer.lock:
                    if outer.speech_failures > 0:
                        outer.speech_failures -= 1
                        return self.refuse(outer.speech_status, "voice busy")
                if outer.speech_delay:
                    time.sleep(outer.speech_delay)
                self.binary(MP3)

            def chat(self, body):
                if body.get("response_format") and outer.refuse_json_mode:
                    return self.refuse(400, "response_format is not supported")
                self.binary(
                    json.dumps({"choices": [{"message": {"content": outer.answer(body)}}]}).encode(),
                    "application/json",
                )

            def refuse(self, status, message):
                data = json.dumps({"error": {"message": message}}).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def binary(self, data, content_type="application/octet-stream"):
                self.send_response(200)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

        self.httpd = HTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def answer(self, body) -> str:
        system = "".join(m.get("content") or "" for m in body.get("messages", [])
                         if m.get("role") == "system")
        user = "".join(m.get("content") or "" for m in body.get("messages", [])
                       if m.get("role") == "user")
        if "instruction for" in system:
            return '{"way": "manner"}'
        if "You map a text" in system:
            found = _TOPICS.search(system)
            return topics_answer(int(found.group(1)) if found else 1)
        if "You plan a podcast episode" in system:
            found = _PARTS.search(system)
            with self.lock:
                files = self.plan_files.pop(0) if self.plan_files else None
            return plan_answer(int(found.group(1)) if found else 1, files)
        found = _COUNT.search(user)
        count = int(found.group(1) or found.group(2)) if found else 4
        return lines_answer(count, self.persons)

    @property
    def base_url(self):
        return "http://127.0.0.1:%d/v1" % self.httpd.server_port

    @property
    def plan_requests(self):
        return [r for r in self.chat_requests
                if "You plan a podcast episode" in str(r["body"].get("messages"))]

    @property
    def topic_requests(self):
        return [r for r in self.chat_requests
                if "You map a text" in str(r["body"].get("messages"))]

    @property
    def chat_requests(self):
        return [r for r in self.requests if r["path"].endswith("/chat/completions")]

    @property
    def speech_requests(self):
        return [r for r in self.requests if r["path"].endswith("/audio/speech")]

    @property
    def spoken(self):
        return [r["body"]["input"] for r in self.speech_requests]

    def stop(self):
        self.httpd.shutdown()
        self.httpd.server_close()


@pytest.fixture
def server():
    s = FakeServer()
    yield s
    s.stop()


@pytest.fixture
def conf(server):
    """What one run needs against the fake server."""
    return {
        "base_url": server.base_url,
        "model": "chat-model",
        "tts_model": "speech-model",
        "voices": "alloy,echo",
        "max_retries": 0,
        "concurrency": 1,
    }


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    import os

    for name in list(os.environ):
        if name.startswith(("DOC2PODCAST_", "DOC2SPEECH_")) or name == "OPENAI_API_KEY":
            monkeypatch.delenv(name)


@pytest.fixture(autouse=True)
def fresh_json_mode():
    """The JSON mode is probed once per process, every test starts with it on."""
    from doc2podcast import client

    client._json_mode_supported = True
    yield
    client._json_mode_supported = True


@pytest.fixture
def material(tmp_path):
    path = tmp_path / "paper.md"
    path.write_text(
        "# The deposit\n\n"
        "A landlord may keep only what the damage costs. The tenant gets the rest back.\n"
        "The bill for heating may be held back for a few months.\n",
        encoding="utf-8",
    )
    return path


def make_pdf(path: Path, lines) -> Path:
    """A minimal PDF with a text layer, one line per entry."""
    text_ops = "BT /F1 12 Tf 72 720 Td " + " ".join(
        "(%s) Tj 0 -16 Td" % line for line in lines
    ) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        "<< /Length %d >>\nstream\n%s\nendstream" % (len(text_ops), text_ops),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += ("%d 0 obj\n%s\nendobj\n" % (number, body)).encode("latin-1")
    xref = len(out)
    out += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)).encode()
    for offset in offsets:
        out += ("%010d 00000 n \n" % offset).encode()
    out += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objects) + 1, xref)).encode()
    path.write_bytes(out)
    return path


def make_docx(path: Path, paragraphs) -> Path:
    from docx import Document

    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(path))
    return path


def make_pptx(path: Path, slides) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[5]
    for title in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
    presentation.save(str(path))
    return path
