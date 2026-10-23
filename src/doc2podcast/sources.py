"""One file into the material of a run, page by page.

What comes out of here is the single truth about a source for the rest of the run.
The text is produced once and never derived again.
"""

import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Sequence, Tuple, Union

from .errors import EmptyTextError, UnsupportedFileError

# What a page break looks like in the assembled text.
_PAGE_BREAK = "\n\n"

# Subtitles are in here because a transcript is a document like any other.
EXTENSIONS: Tuple[str, ...] = (".pdf", ".txt", ".md", ".docx", ".pptx", ".srt", ".vtt")

Content = Union[bytes, str, Path]


# Puts standalone accents from LaTeX PDFs back on their letter.
_ACCENT_SCAN_RANGE = range(0x00A0, 0x0300)

# The acute stays out, it is typed as an apostrophe.
_ACCENT_EXCLUDED = {"´"}

# Circumflex and caron added by hand, Unicode knows no decomposition for them.
_ACCENT_EXTRA = {
    "ˆ": "̂",  # MODIFIER LETTER CIRCUMFLEX ACCENT -> COMBINING CIRCUMFLEX
    "ˇ": "̌",  # CARON -> COMBINING CARON
}


def _spacing_accent_marks():
    """Standalone accent to combining accent, derived from the Unicode data."""
    marks = {}
    for code_point in _ACCENT_SCAN_RANGE:
        char = chr(code_point)
        if char in _ACCENT_EXCLUDED:
            continue
        parts = unicodedata.decomposition(char).split()
        if len(parts) == 3 and parts[0] == "<compat>" and parts[1] == "0020":
            marks[char] = chr(int(parts[2], 16))
    marks.update(_ACCENT_EXTRA)
    return marks


_SPACING_ACCENTS = _spacing_accent_marks()
_SPACING_ACCENT_RE = re.compile("[%s]" % re.escape("".join(_SPACING_ACCENTS)))


def _attach_spacing_accents(text: str) -> str:
    """Attach every standalone accent to the letter it belongs to."""
    out: List[str] = []
    index = 0
    end = len(text)
    while index < end:
        char = text[index]
        mark = _SPACING_ACCENTS.get(char)
        if mark:
            # The following letter takes precedence.
            if index + 1 < end and text[index + 1].isalpha():
                out.append(text[index + 1] + mark)
                index += 2
                continue
            # Otherwise the preceding one, as in "A¯" or "p(A¯)".
            if out and out[-1][-1:].isalpha():
                out[-1] += mark
                index += 1
                continue
            # No letter next to it means it cannot be assigned, then it stays.
        out.append(char)
        index += 1
    return "".join(out)


def clean(text: str) -> str:
    """Removes null bytes and reassembles separated accents."""
    if not text:
        return text
    text = text.replace("\x00", "")
    if _SPACING_ACCENT_RE.search(text):
        text = _attach_spacing_accents(text)
    return unicodedata.normalize("NFC", text)


@dataclass(frozen=True)
class Source:
    """A read file as the rest of the run sees it."""

    name: str
    text: str
    pages: int
    # Where every page starts and ends in ``text``: (number, from, to).
    bounds: Tuple[Tuple[int, int, int], ...] = ()


def _as_bytes(content: Content) -> bytes:
    if isinstance(content, bytes):
        return content
    with open(content, "rb") as handle:
        return handle.read()


def _as_stream(content: Content):
    """What python-docx and python-pptx should open: a path stays a path."""
    return BytesIO(content) if isinstance(content, bytes) else str(content)


def _pdf_pages(content: Content) -> List[str]:
    """The text layer of every page. A scanned PDF without a text layer yields nothing."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(content if isinstance(content, bytes) else str(content))
    try:
        pages: List[str] = []
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    raw = textpage.get_text_range() or ""
                finally:
                    textpage.close()
            finally:
                page.close()
            pages.append(clean(raw.replace("\r\n", "\n").replace("\r", "\n")).strip())
        return pages
    finally:
        pdf.close()


def _pptx_slides(content: Content) -> List[str]:
    """The text of every slide on its own, without footer, slide number and date."""
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    # Footer, slide number and date stand on every slide and say nothing about it.
    furniture = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}
    presentation = Presentation(_as_stream(content))
    slides: List[str] = []
    for slide in presentation.slides:
        texts: List[str] = []
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.type in furniture:
                continue
            _collect_shape_text(shape, texts)
        slides.append(clean("\n".join(texts)).strip())
    return slides


def _collect_shape_text(shape, texts: List[str]) -> None:
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                texts.append(text)
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)


def _docx_page(content: Content) -> str:
    from docx import Document

    document = Document(_as_stream(content))
    paragraphs = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)
    return clean("\n".join(paragraphs)).strip()


_MD_CODE_BLOCK = re.compile(r"```[\s\S]*?```")
_MD_CODE_SPAN = re.compile(r"`[^`]+`")
_MD_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_MD_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_MD_HEADING = re.compile(r"^#{1,6}\s+", re.MULTILINE)
_MD_BOLD = re.compile(r"\*{1,3}([^*]+)\*{1,3}")
_MD_ITALIC = re.compile(r"_{1,3}([^_]+)_{1,3}")
_MD_RULE = re.compile(r"^[-*_]{3,}\s*$", re.MULTILINE)
_MD_TAG = re.compile(r"<[^>]+>")
_BLANK_LINES = re.compile(r"\n{3,}")


def _md_page(content: Content) -> str:
    text = _text_page(content)
    text = _MD_CODE_BLOCK.sub("", text)
    text = _MD_CODE_SPAN.sub("", text)
    text = _MD_IMAGE.sub("", text)
    text = _MD_LINK.sub(r"\1", text)
    text = _MD_HEADING.sub("", text)
    text = _MD_BOLD.sub(r"\1", text)
    text = _MD_ITALIC.sub(r"\1", text)
    text = _MD_RULE.sub("", text)
    text = _MD_TAG.sub("", text)
    return _BLANK_LINES.sub("\n\n", text).strip()


def _text_page(content: Content) -> str:
    raw = _as_bytes(content).decode("utf-8", errors="replace")
    return clean(raw.replace("\r\n", "\n").replace("\r", "\n")).strip()


_SRT_INDEX = re.compile(r"^\d+$")
_SRT_TIME = re.compile(r"\d{1,2}:\d{2}:\d{2}[.,]\d{1,3}\s*-->")


def _subtitle_page(content: Content) -> str:
    """A transcript without its numbers and timings, every line said only once."""
    lines: List[str] = []
    for raw in _text_page(content).splitlines():
        line = raw.strip()
        if not line or line.upper().startswith("WEBVTT") or line.startswith("NOTE"):
            continue
        if _SRT_INDEX.match(line) or _SRT_TIME.search(line):
            continue
        if lines and lines[-1] == line:
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def _pages(name: str, content: Content) -> List[str]:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return _pdf_pages(content)
    if suffix == ".pptx":
        # Every slide is a page of its own, the way a PDF page is.
        return _pptx_slides(content)
    if suffix == ".docx":
        return [_docx_page(content)]
    if suffix == ".md":
        return [_md_page(content)]
    if suffix == ".txt":
        return [_text_page(content)]
    if suffix in (".srt", ".vtt"):
        return [_subtitle_page(content)]
    raise UnsupportedFileError(
        "File type %s is not supported. Allowed: %s"
        % (suffix or "(no extension)", ", ".join(EXTENSIONS))
    )


def _assemble(pages: Sequence[str]) -> Tuple[str, List[Tuple[int, int, int]]]:
    """The pages as one text, plus where each page starts and ends in it.

    An empty page keeps its number but takes no room, so the page of a passage stays
    right even in a document with blank pages in the middle.
    """
    parts: List[str] = []
    bounds: List[Tuple[int, int, int]] = []
    at = 0
    for number, raw in enumerate(pages, start=1):
        # Stripped, otherwise leading whitespace would count as page content.
        page = raw.strip()
        if not page:
            continue
        if parts:
            at += len(_PAGE_BREAK)
        bounds.append((number, at, at + len(page)))
        parts.append(page)
        at += len(page)
    return _PAGE_BREAK.join(parts), bounds


def read_bytes(name: str, data: bytes) -> Source:
    """An uploaded file to a source record, the type comes from its name."""
    return _source(name or "the file", data)


def read_file(path) -> Source:
    """One file to a source record. Raises when it carries no text."""
    path = Path(path)
    if not path.is_file():
        raise UnsupportedFileError("File not found: %s" % path)
    return _source(path.name, path)


def _source(name: str, content: Content) -> Source:
    try:
        pages = _pages(name, content)
    except UnsupportedFileError:
        raise
    except UnicodeDecodeError:
        raise UnsupportedFileError("%s is not UTF-8 text." % name)
    except Exception as exc:
        raise UnsupportedFileError("%s cannot be read: %s" % (name, exc))
    text, bounds = _assemble(pages)
    if not text.strip():
        raise EmptyTextError("No text found in %s." % name)
    return Source(name=name, text=text, pages=len(pages), bounds=tuple(bounds))


def read_all(paths) -> List[Source]:
    """Every given file as a source, in the order they were named."""
    return [read_file(path) for path in paths]


def material(sources: Sequence[Source], max_chars: int = 0) -> str:
    """The material one episode is written from, every page naming its file and page.

    The square brackets are the reason the prompt tells the model never to read them
    out: they say where a sentence comes from, nothing more. When the material is
    longer than ``max_chars``, every source keeps the same share of what fits.
    """
    if not sources:
        raise EmptyTextError("No material to work with.")
    budget = _budget(sources, max_chars)
    blocks: List[str] = []
    for source in sources:
        text = source.text[: budget[source.name]] if budget else source.text
        for number, start, end in source.bounds or ((1, 0, len(text)),):
            piece = text[start:end].strip()
            if not piece:
                continue
            blocks.append("[%s, p. %d]\n%s" % (source.name, number, piece))
    if not blocks:
        raise EmptyTextError("No material to work with.")
    return "\n\n".join(blocks)


def _budget(sources: Sequence[Source], max_chars: int) -> Dict[str, int]:
    """How many characters each source may contribute, empty when everything fits."""
    total = sum(len(s.text) for s in sources)
    if not max_chars or total <= max_chars:
        return {}
    share = max_chars / float(total)
    return {s.name: max(1, int(len(s.text) * share)) for s in sources}
